"""Symmetric reader and structural validator for generated import workbooks."""

from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import load_workbook

from .excel_writer import (
    HEADER_INDEX,
    STOCKPILE_HEADERS,
    STOCKPILE_SCHEMA_VERSION,
    STOCKPILE_SHEET,
    TEMPLATE_HEADERS,
    _fingerprint,
)


class WorkbookValidationError(ValueError):
    def __init__(self, errors: list[dict[str, Any]]):
        self.errors = errors
        super().__init__("; ".join(str(error.get("message") or "Invalid workbook") for error in errors))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _bool(value: Any) -> bool | None:
    normalized = _text(value).casefold()
    if normalized in {"true", "1", "yes", "on"}:
        return True
    if normalized in {"false", "0", "no", "off"}:
        return False
    return None


def _number(value: Any, *, integer: bool = False) -> int | float | None:
    if value in (None, ""):
        return None
    try:
        return int(value) if integer else float(value)
    except (TypeError, ValueError):
        return None


def _identity(row: list[Any]) -> dict[str, str]:
    return {
        "handle": _text(row[HEADER_INDEX["Handle"]]).casefold(),
        "option1": _text(row[HEADER_INDEX["Option1 Value"]]),
        "option2": _text(row[HEADER_INDEX["Option2 Value"]]),
        "option3": _text(row[HEADER_INDEX["Option3 Value"]]),
        "sku": _text(row[HEADER_INDEX["Variant SKU"]]),
    }


def _variant(row: list[Any]) -> dict[str, Any]:
    fields: dict[str, tuple[str, str]] = {
        "option1": ("Option1 Value", "text"),
        "option2": ("Option2 Value", "text"),
        "option3": ("Option3 Value", "text"),
        "sku": ("Variant SKU", "text"),
        "grams": ("Variant Grams", "int"),
        "inventory_tracker": ("Variant Inventory Tracker", "text"),
        "inventory_quantity": ("Variant Inventory Qty", "int"),
        "inventory_policy": ("Variant Inventory Policy", "text"),
        "fulfillment_service": ("Variant Fulfillment Service", "text"),
        "price": ("Variant Price", "number"),
        "compare_at_price": ("Variant Compare At Price", "number"),
        "requires_shipping": ("Variant Requires Shipping", "bool"),
        "taxable": ("Variant Taxable", "bool"),
        "barcode": ("Variant Barcode", "text"),
        "variant_image": ("Variant Image", "text"),
        "variant_weight_unit": ("Variant Weight Unit", "text"),
        "variant_tax_code": ("Variant Tax Code", "text"),
        "cost_per_item": ("Cost per item", "number"),
        "price_international": ("Price / International", "text"),
        "compare_at_price_international": ("Compare At Price / International", "text"),
    }
    result: dict[str, Any] = {}
    for key, (header, kind) in fields.items():
        raw = row[HEADER_INDEX[header]]
        value: Any = _text(raw)
        if kind == "int":
            value = _number(raw, integer=True)
        elif kind == "number":
            value = _number(raw)
        elif kind == "bool":
            value = _bool(raw)
        if value not in (None, ""):
            result[key] = value
    return result


def _product(first: list[Any], rows: list[list[Any]], origins: dict[str, Any]) -> dict[str, Any]:
    text_fields = {
        "handle": "Handle",
        "title": "Title",
        "body_html": "Body (HTML)",
        "vendor": "Vendor",
        "product_category": "Product Category",
        "product_type": "Type",
        "seo_title": "SEO Title",
        "seo_description": "SEO Description",
        "status": "Status",
        "image_src": "Image Src",
        "image_alt_text": "Image Alt Text",
    }
    product: dict[str, Any] = {
        key: _text(first[HEADER_INDEX[header]])
        for key, header in text_fields.items()
        if _text(first[HEADER_INDEX[header]])
    }
    tags = _text(first[HEADER_INDEX["Tags"]])
    if tags:
        product["tags"] = [tag.strip() for tag in tags.split(",") if tag.strip()]
    for key, header in (("published", "Published"), ("gift_card", "Gift Card")):
        value = _bool(first[HEADER_INDEX[header]])
        if value is not None:
            product[key] = value
    image_position = _number(first[HEADER_INDEX["Image Position"]], integer=True)
    if image_position is not None:
        product["image_position"] = image_position
    option_names = [
        _text(first[HEADER_INDEX[f"Option{index} Name"]])
        for index in (1, 2, 3)
    ]
    options = []
    for index, name in enumerate(option_names, start=1):
        if not name:
            continue
        values = []
        for row in rows:
            value = _text(row[HEADER_INDEX[f"Option{index} Value"]])
            if value and value not in values:
                values.append(value)
        options.append({"name": name, "values": values})
    if options:
        product["options"] = options
    product["variants"] = [_variant(row) for row in rows]
    product["value_origins"] = origins
    return product


def read_generated_workbook(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    errors: list[dict[str, Any]] = []
    if "Products" not in workbook.sheetnames or STOCKPILE_SHEET not in workbook.sheetnames:
        raise WorkbookValidationError([{"cell": "Products!A1", "message": "Generated workbook metadata is missing"}])
    products_sheet = workbook["Products"]
    metadata_sheet = workbook[STOCKPILE_SHEET]
    headers = [_text(cell.value) for cell in products_sheet[1]]
    if headers != TEMPLATE_HEADERS:
        raise WorkbookValidationError([{"cell": "Products!A1", "message": "Products headers were changed"}])
    metadata_headers = [_text(cell.value) for cell in metadata_sheet[1]]
    if metadata_headers != STOCKPILE_HEADERS:
        raise WorkbookValidationError([{"cell": f"{STOCKPILE_SHEET}!A1", "message": "Workbook metadata headers are invalid"}])

    visible = [list(row) for row in products_sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]
    metadata = [
        dict(zip(STOCKPILE_HEADERS, row))
        for row in metadata_sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]
    if len(visible) != len(metadata) or products_sheet.max_row != len(metadata) + 1:
        raise WorkbookValidationError([{"cell": "Products!A2", "message": "Product or variant rows were added or deleted"}])

    expected_by_product: dict[str, list[str]] = {}
    current_by_product: dict[str, list[str]] = {}
    for index, (row, meta) in enumerate(zip(visible, metadata), start=2):
        if int(meta.get("schema_version") or 0) != STOCKPILE_SCHEMA_VERSION:
            errors.append({"cell": f"{STOCKPILE_SHEET}!A{index}", "message": "Unsupported workbook schema version"})
        if int(meta.get("visible_row") or 0) != index:
            errors.append({"cell": f"Products!A{index}", "message": "Product or variant rows were reordered"})
        if _text(row[HEADER_INDEX["Handle"]]).casefold() != _text(meta.get("handle")).casefold():
            errors.append({"cell": f"Products!A{index}", "message": "Product identity cannot be changed"})
        product_id = _text(meta.get("product_row_id"))
        expected_by_product.setdefault(product_id, []).append(_text(meta.get("identity_fingerprint")))
        current_by_product.setdefault(product_id, []).append(_fingerprint(_identity(row)))
    for product_id, expected in expected_by_product.items():
        current = current_by_product[product_id]
        for offset, (before, after) in enumerate(zip(expected, current)):
            if before != after and after in expected and expected.index(after) != offset:
                errors.append({"cell": f"Products!A{offset + 2}", "message": "Variant rows were reordered"})
    if errors:
        raise WorkbookValidationError(errors)

    groups: list[tuple[list[list[Any]], list[dict[str, Any]]]] = []
    for row, meta in zip(visible, metadata):
        product_id = _text(meta.get("product_row_id"))
        if not groups or _text(groups[-1][1][0].get("product_row_id")) != product_id:
            groups.append(([], []))
        groups[-1][0].append(row)
        groups[-1][1].append(meta)

    products: list[dict[str, Any]] = []
    for rows, metas in groups:
        try:
            origins = json.loads(_text(metas[0].get("value_origins")) or "{}")
        except json.JSONDecodeError:
            origins = {}
        products.append(_product(rows[0], rows, origins if isinstance(origins, dict) else {}))

    if "AI Enhancements" in workbook.sheetnames:
        ai_sheet = workbook["AI Enhancements"]
        for row in ai_sheet.iter_rows(min_row=2, values_only=True):
            handle, _, attribute, value, value_type = list(row)[:5]
            attribute = _text(attribute)
            if not attribute.startswith("metafield:"):
                continue
            target = next((product for product in products if product.get("handle") == _text(handle)), None)
            if target is None:
                continue
            namespace_key = attribute.split(":", 1)[1]
            namespace, _, key = namespace_key.partition(".")
            target.setdefault("metafields", []).append({
                "namespace": namespace,
                "key": key,
                "value": _text(value),
                "type": _text(value_type) or "single_line_text_field",
            })
    workbook.close()
    return products
