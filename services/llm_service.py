from __future__ import annotations

import base64
import csv
from functools import wraps
from io import BytesIO, StringIO
import json
import logging
import os
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from time import perf_counter
from typing import Any, Callable

from agent_framework.github import GitHubCopilotAgent
from agent_framework_monty import FileMount, MontyCodeActProvider
from copilot import CopilotClient, ToolSet
from copilot.session import (
    PermissionDecisionApproveOnce,
    PermissionDecisionUserNotAvailable,
)
from copilot.tools import Tool, ToolBinaryResult, ToolInvocation, ToolResult
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from ai.models import ProductsList
from application.ports.llm_port import ExtractionRequest, ModelConfiguration
from shared.metrics_signals import signal_llm_latency
from shared.token_usage import normalize_reported_token_usage
from .interfaces import CollaboraServiceInterface, SupabaseServiceInterface


logger = logging.getLogger(__name__)
MAX_EVIDENCE_PAGE_SIZE = 200
MAX_RANGE_CELLS = 2_000
MAX_PREVIEW_PAGES = 12
LOCAL_COLLABORA_URL = "http://localhost:8080"
MONTY_RESOURCE_LIMITS = {
    "max_duration_secs": 30.0,
    "max_memory": 128 * 1024 * 1024,
}


def _strip_markdown_json_fence(text: str) -> str:
    match = re.match(r"^```(?:json)?\s*([\s\S]*?)\s*```$", text.strip(), re.IGNORECASE)
    return match.group(1).strip() if match else text.strip()


def sanitize_source_filename(filename: str) -> str:
    """Return a safe, non-hidden basename for the one staged source file."""
    basename = Path(filename.replace("\\", "/")).name.strip()
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", basename).strip("._")
    return safe or "source.bin"


def resolve_collabora_url(request: ExtractionRequest) -> tuple[str, str]:
    """Resolve the conversion endpoint without ever tracing the endpoint value."""
    if request.collabora_base_url and request.collabora_base_url.strip():
        return request.collabora_base_url.strip(), "job"
    configured = os.getenv("COLLABORA_URL", "").strip()
    if configured:
        return configured, "worker_env"
    return LOCAL_COLLABORA_URL, "local_default"


def map_copilot_provider(model: ModelConfiguration) -> dict[str, Any]:
    """Map stored OpenAI/OpenAI-compatible BYOK settings to the Copilot SDK."""
    provider = model.provider.strip().casefold()
    if not model.api_key.strip():
        raise RuntimeError(
            "COPILOT_BYOK_AUTH_REQUIRED: active model API key is missing"
        )
    if not model.model_id.strip():
        raise RuntimeError("COPILOT_BYOK_MODEL_REQUIRED: active model id is missing")
    if not model.base_url.strip():
        raise RuntimeError(
            "COPILOT_BYOK_ENDPOINT_REQUIRED: active model endpoint is missing"
        )
    if not (
        provider.startswith("openai")
        or "openai-compat" in provider
        or "ollama" in provider
    ):
        raise RuntimeError(f"COPILOT_BYOK_PROVIDER_UNSUPPORTED: {model.provider}")
    mapped: dict[str, Any] = {
        "type": "openai",
        "wire_api": (
            "responses"
            if "api.openai.com" in model.base_url.casefold()
            else "completions"
        ),
        "base_url": model.base_url.rstrip("/"),
        "api_key": model.api_key,
        "model_id": model.model_id,
        "wire_model": model.model_id,
    }
    if model.max_tokens is not None:
        mapped["max_output_tokens"] = model.max_tokens
    return mapped


def scoped_permission_handler(
    allowed_custom_tools: set[str] | frozenset[str],
    trace_event: Callable[..., Any] | None = None,
) -> Callable[[Any, dict[str, str]], Any]:
    """Approve only explicitly registered custom tools; deny every other capability."""
    allowed = frozenset(allowed_custom_tools)

    def handle(request: Any, _invocation: dict[str, str]) -> Any:
        kind = str(getattr(request, "kind", "") or "")
        tool_name = str(getattr(request, "tool_name", "") or "")
        approved = kind == "custom-tool" and tool_name in allowed
        decision = (
            PermissionDecisionApproveOnce()
            if approved
            else PermissionDecisionUserNotAvailable()
        )
        if trace_event:
            trace_event(
                phase="copilot_permission_decision",
                message=(
                    "Approved scoped Copilot tool"
                    if approved
                    else "Denied Copilot capability"
                ),
                payload_preview={
                    "kind": kind or "unknown",
                    "tool": tool_name or None,
                    "decision": decision.kind,
                },
            )
        return decision

    return handle


deny_permission = scoped_permission_handler(frozenset())


def _trace_tool(
    tool: Callable[..., Any],
    *,
    name: str,
    trace_event: Callable[..., Any] | None,
) -> Callable[..., Any]:
    """Trace tool lifecycle without retaining arguments, results, or exception text."""
    if trace_event is None:
        return tool

    @wraps(tool)
    async def traced(*args: Any, **kwargs: Any) -> Any:
        trace_event(
            phase="copilot_tool_start",
            message=f"Started Copilot tool {name}",
            payload_preview={"tool": name},
        )
        try:
            result = await tool(*args, **kwargs)
        except Exception as exc:
            trace_event(
                phase="copilot_tool_failure",
                message=f"Copilot tool {name} failed",
                level="error",
                payload_preview={"tool": name, "error_type": type(exc).__name__},
            )
            raise
        trace_event(
            phase="copilot_tool_success",
            message=f"Completed Copilot tool {name}",
            payload_preview={"tool": name},
        )
        return result

    return traced


class RunScopedExtractionTools:
    def __init__(self, request: ExtractionRequest) -> None:
        self.request = request
        self.submitted: ProductsList | None = None
        self._sheets: dict[str, list[list[Any]]] | None = None

    async def document_manifest(self) -> dict[str, Any]:
        """Return source kind, page/sheet names, and evidence counts without source bytes."""
        if self.request.evidence_catalog is not None:
            pages = sorted(
                {entry.page for entry in self.request.evidence_catalog.entries}
            )
            return {
                "kind": self.request.document_kind,
                "pages": pages,
                "evidence_count": len(self.request.evidence_catalog.entries),
            }
        sheets = self._load_sheets()
        return {
            "kind": self.request.document_kind,
            "sheets": [
                {
                    "name": name,
                    "rows": len(rows),
                    "columns": max((len(row) for row in rows), default=0),
                }
                for name, rows in sheets.items()
            ],
        }

    async def read_evidence_page(
        self, page: int = 1, page_size: int = 100
    ) -> dict[str, Any]:
        """Read one bounded, one-based page of the canonical PDF evidence catalog."""
        catalog = self.request.evidence_catalog
        if catalog is None:
            raise ValueError("Evidence catalog is unavailable for this document")
        size = max(1, min(int(page_size), MAX_EVIDENCE_PAGE_SIZE))
        number = max(1, int(page))
        start = (number - 1) * size
        entries = catalog.entries[start : start + size]
        return {
            "page": number,
            "page_size": size,
            "total": len(catalog.entries),
            "items": [
                {"evidence_index": item.evidence_index, "text": item.text}
                for item in entries
            ],
        }

    async def search_evidence(
        self, query: str, page: int = 1, page_size: int = 50
    ) -> dict[str, Any]:
        """Search canonical PDF evidence text and return bounded paginated matches."""
        catalog = self.request.evidence_catalog
        if catalog is None:
            raise ValueError("Evidence catalog is unavailable for this document")
        needle = query.strip().casefold()
        if not needle:
            raise ValueError("query must not be empty")
        matches = [item for item in catalog.entries if needle in item.text.casefold()]
        size = max(1, min(int(page_size), MAX_EVIDENCE_PAGE_SIZE))
        number = max(1, int(page))
        start = (number - 1) * size
        return {
            "page": number,
            "page_size": size,
            "total": len(matches),
            "items": [
                {"evidence_index": item.evidence_index, "text": item.text}
                for item in matches[start : start + size]
            ],
        }

    async def read_workbook_range(
        self,
        sheet: str,
        start_row: int,
        end_row: int,
        start_column: int = 1,
        end_column: int = 20,
    ) -> dict[str, Any]:
        """Read a capped worksheet/CSV range with deterministic A1 cell references."""
        sheets = self._load_sheets()
        if sheet not in sheets:
            raise ValueError(f"Unknown sheet: {sheet}")
        sr, er = max(1, int(start_row)), max(1, int(end_row))
        sc, ec = max(1, int(start_column)), max(1, int(end_column))
        if er < sr or ec < sc:
            raise ValueError("Range end must not precede range start")
        if (er - sr + 1) * (ec - sc + 1) > MAX_RANGE_CELLS:
            raise ValueError(f"Range exceeds {MAX_RANGE_CELLS} cells")
        rows = sheets[sheet]
        cells = []
        for row_no in range(sr, er + 1):
            row = rows[row_no - 1] if row_no <= len(rows) else []
            values = []
            for col_no in range(sc, ec + 1):
                value = row[col_no - 1] if col_no <= len(row) else None
                values.append(
                    {"cell": f"{get_column_letter(col_no)}{row_no}", "value": value}
                )
            cells.append(values)
        return {
            "sheet": sheet,
            "range": f"{get_column_letter(sc)}{sr}:{get_column_letter(ec)}{er}",
            "rows": cells,
        }

    async def validate_extraction(self, payload: dict[str, Any]) -> dict[str, Any]:
        """Validate a complete Shopify extraction payload without submitting it."""
        try:
            parsed = ProductsList.model_validate(payload)
        except ValidationError as exc:
            return {"valid": False, "errors": exc.errors(include_url=False)}
        return {"valid": True, "products": len(parsed.products)}

    async def submit_extraction(self, payload: dict[str, Any]) -> dict[str, Any]:
        """Validate and submit the one complete extraction payload returned to the server."""
        if self.submitted is not None:
            raise RuntimeError("Extraction was already submitted for this run")
        self.submitted = ProductsList.model_validate(payload)
        return {"accepted": True, "products": len(self.submitted.products)}

    def _load_sheets(self) -> dict[str, list[list[Any]]]:
        if self._sheets is not None:
            return self._sheets
        if self.request.document_kind == "csv":
            rows = list(
                csv.reader(
                    StringIO(self.request.content.decode("utf-8-sig", errors="replace"))
                )
            )
            self._sheets = {"Sheet1": rows}
            return self._sheets
        if self.request.document_kind != "spreadsheet":
            raise ValueError("Workbook ranges are unavailable for this document")
        workbook = load_workbook(
            BytesIO(self.request.content), read_only=True, data_only=True
        )
        try:
            self._sheets = {
                sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()
        return self._sheets


class PagePreviewTool:
    def __init__(
        self,
        request: ExtractionRequest,
        collabora: CollaboraServiceInterface,
        directory: Path,
        *,
        collabora_url: str | None = None,
        trace_event: Callable[..., Any] | None = None,
    ) -> None:
        self.request = request
        self.collabora = collabora
        self.directory = directory
        self.rendered_pages: set[int] = set()
        self._pdf: bytes | None = (
            request.content if request.document_kind == "pdf" else None
        )
        self._collabora_url = collabora_url or resolve_collabora_url(request)[0]
        self._trace_event = trace_event
        self._preview_unavailable = False

    async def view_page(self, page_number: int) -> ToolResult:
        """Render and view one source page on demand; at most 12 unique pages per extraction."""
        if self.request.document_kind == "csv":
            raise ValueError("CSV sources do not have visual pages")
        number = int(page_number)
        if number < 1:
            raise ValueError("page_number is one-based")
        if (
            number not in self.rendered_pages
            and len(self.rendered_pages) >= MAX_PREVIEW_PAGES
        ):
            raise ValueError(
                f"Page preview limit of {MAX_PREVIEW_PAGES} unique pages reached"
            )
        if self._preview_unavailable:
            return ToolResult(
                text_result_for_llm=(
                    "Visual preview is unavailable for this extraction. Use workbook ranges or PDF evidence instead."
                )
            )
        if self._pdf is None:
            try:
                self._pdf = await self.collabora.convert_document_to_pdf_collabora(
                    self.request.content,
                    filename=self.request.filename,
                    content_type=self.request.content_type,
                    collabora_base_url=self._collabora_url,
                )
            except Exception as exc:
                self._preview_unavailable = True
                if self._trace_event:
                    self._trace_event(
                        phase="copilot_preview_unavailable",
                        message="Collabora preview conversion failed; range extraction remains available",
                        level="error",
                        payload_preview={"error_type": type(exc).__name__},
                    )
                    self._trace_event(
                        phase="copilot_preview_fallback",
                        message="Falling back to workbook ranges or PDF evidence",
                        payload_preview={"preview_available": False},
                    )
                return ToolResult(
                    text_result_for_llm=(
                        "Visual preview conversion failed and is disabled for this extraction. "
                        "Use workbook ranges or PDF evidence instead."
                    )
                )
        import fitz

        document = fitz.open(stream=self._pdf, filetype="pdf")
        try:
            if number > document.page_count:
                raise ValueError(
                    f"page_number exceeds document page count ({document.page_count})"
                )
            output = self.directory / f"page-{number}.png"
            if not output.exists():
                output.write_bytes(
                    document[number - 1]
                    .get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    .tobytes("png")
                )
            data = output.read_bytes()
        finally:
            document.close()
        self.rendered_pages.add(number)
        return ToolResult(
            text_result_for_llm=f"Rendered source page {number}.",
            binary_results_for_llm=[
                ToolBinaryResult(
                    data=base64.b64encode(data).decode("ascii"),
                    mime_type="image/png",
                    type="image",
                    description=f"Source page {number}",
                )
            ],
        )

    def as_copilot_tool(self) -> Tool:
        """Expose preview natively so the SDK preserves binary image content."""

        async def handler(invocation: ToolInvocation) -> ToolResult:
            arguments = invocation.arguments or {}
            if self._trace_event:
                self._trace_event(
                    phase="copilot_tool_start",
                    message="Started Copilot tool view_page",
                    payload_preview={"tool": "view_page"},
                )
            try:
                result = await self.view_page(arguments.get("page_number"))
            except Exception as exc:
                if self._trace_event:
                    self._trace_event(
                        phase="copilot_tool_failure",
                        message="Copilot tool view_page failed",
                        level="error",
                        payload_preview={
                            "tool": "view_page",
                            "error_type": type(exc).__name__,
                        },
                    )
                raise
            if self._trace_event:
                self._trace_event(
                    phase="copilot_tool_success",
                    message="Completed Copilot tool view_page",
                    payload_preview={"tool": "view_page"},
                )
            return result

        tool = Tool(
            name="view_page",
            description=self.view_page.__doc__ or "Render one source page.",
            handler=handler,
            parameters={
                "type": "object",
                "properties": {
                    "page_number": {
                        "type": "integer",
                        "minimum": 1,
                        "description": "One-based source page number.",
                    }
                },
                "required": ["page_number"],
                "additionalProperties": False,
            },
        )
        # Retained only on the in-process tool object; useful for lifecycle tests
        # and never serialized or sent to the model.
        tool._preview_directory = self.directory  # type: ignore[attr-defined]
        return tool


class LLMService:
    def __init__(
        self,
        *,
        collabora: CollaboraServiceInterface,
        supabase: SupabaseServiceInterface,
        agent_factory=GitHubCopilotAgent,
    ) -> None:
        self.collabora = collabora
        self.supabase = supabase
        self.agent_factory = agent_factory

    async def extract(self, request: ExtractionRequest) -> ProductsList:
        started_at = perf_counter()
        status = "failed"
        try:
            prepared = await self._prepare_legacy_spreadsheet(request)
            result = await self._run_copilot(prepared)
            status = "succeeded"
            return result
        finally:
            signal_llm_latency(
                operation="copilot_extract",
                status=status,
                duration_ms=int((perf_counter() - started_at) * 1000),
                model_provider=request.model.provider,
                model_id=request.model.model_id,
            )

    async def _prepare_legacy_spreadsheet(
        self, request: ExtractionRequest
    ) -> ExtractionRequest:
        if request.document_kind != "spreadsheet_legacy":
            return request
        collabora_url, collabora_source = resolve_collabora_url(request)
        if request.trace_event:
            request.trace_event(
                phase="copilot_collabora_resolution",
                message="Resolved Collabora legacy conversion endpoint source",
                payload_preview={"source": collabora_source},
            )
        converted = await self.collabora.convert_document_to_xlsx_collabora(
            request.content,
            filename=request.filename,
            content_type=request.content_type,
            collabora_base_url=collabora_url,
        )
        return ExtractionRequest(
            document_kind="spreadsheet",
            content=converted,
            filename=f"{Path(request.filename).stem}.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            model=request.model,
            evidence_catalog=request.evidence_catalog,
            validation_feedback=request.validation_feedback,
            extraction_mode=request.extraction_mode,
            collabora_base_url=request.collabora_base_url,
            trace_event=request.trace_event,
        )

    async def _run_copilot(self, request: ExtractionRequest) -> ProductsList:
        trace = request.trace_event
        if trace:
            trace(
                phase="copilot_session_start",
                message="Starting isolated Copilot BYOK extraction",
                payload_preview={
                    "provider": request.model.provider,
                    "model": request.model.model_id,
                    "kind": request.document_kind,
                },
            )
        scoped = RunScopedExtractionTools(request)
        all_scoped_tools = [
            _trace_tool(
                scoped.document_manifest, name="document_manifest", trace_event=trace
            ),
            _trace_tool(
                scoped.read_evidence_page, name="read_evidence_page", trace_event=trace
            ),
            _trace_tool(
                scoped.search_evidence, name="search_evidence", trace_event=trace
            ),
            _trace_tool(
                scoped.read_workbook_range,
                name="read_workbook_range",
                trace_event=trace,
            ),
            _trace_tool(
                scoped.validate_extraction,
                name="validate_extraction",
                trace_event=trace,
            ),
            _trace_tool(
                scoped.submit_extraction, name="submit_extraction", trace_event=trace
            ),
        ]
        if request.document_kind == "pdf":
            direct_tools = [all_scoped_tools[index] for index in (0, 1, 2, 4, 5)]
        else:
            direct_tools = [all_scoped_tools[index] for index in (0, 3, 4, 5)]
        with (
            TemporaryDirectory(prefix="stockpile-source-") as source_dir,
            TemporaryDirectory(prefix="stockpile-copilot-state-") as state_dir,
            TemporaryDirectory(prefix="stockpile-preview-") as preview_dir,
        ):
            source_path = Path(source_dir) / sanitize_source_filename(request.filename)
            source_path.write_bytes(request.content)
            source_path.chmod(0o444)
            collabora_url, collabora_source = resolve_collabora_url(request)
            if trace:
                trace(
                    phase="copilot_collabora_resolution",
                    message="Resolved Collabora preview endpoint source",
                    payload_preview={"source": collabora_source},
                )
            monty = MontyCodeActProvider(
                tools=direct_tools,
                file_mounts=[FileMount(source_dir, "/input", "read-only")],
                resource_limits=MONTY_RESOURCE_LIMITS,
            )
            preview = PagePreviewTool(
                request,
                self.collabora,
                Path(preview_dir),
                collabora_url=collabora_url,
                trace_event=trace,
            )
            copilot_tools: list[Any] = list(direct_tools)
            if request.document_kind != "csv":
                copilot_tools.append(preview.as_copilot_tool())
            allowed_custom_tools = {
                "execute_code",
                *(tool.name for tool in monty.get_tools()),
            }
            if request.document_kind != "csv":
                allowed_custom_tools.add("view_page")
            available = ToolSet()
            for tool_name in sorted(allowed_custom_tools):
                available.add_custom(tool_name)
            client = CopilotClient(
                mode="copilot-cli",
                working_directory=source_dir,
                base_directory=state_dir,
                use_logged_in_user=False,
                env={"COPILOT_DISABLE_KEYTAR": "1"},
                telemetry=None,
                enable_remote_sessions=False,
            )
            prompt = self._prompt(request)
            options = {
                "model": request.model.model_id,
                "provider": map_copilot_provider(request.model),
                "timeout": request.model.timeout_seconds or 120.0,
                "on_permission_request": scoped_permission_handler(
                    allowed_custom_tools, trace
                ),
                "available_tools": available,
                "working_directory": source_dir,
                "enable_session_telemetry": False,
                "skip_custom_instructions": True,
                "custom_agents_local_only": True,
                "coauthor_enabled": False,
                "manage_schedule_enabled": False,
                "mcp_servers": {},
                "mcp_oauth_token_storage": "in-memory",
                "embedding_cache_storage": "in-memory",
                "custom_agents": [],
                "enable_config_discovery": False,
                "skip_embedding_retrieval": True,
                "organization_custom_instructions": "",
                "enable_on_demand_instruction_discovery": False,
                "enable_file_hooks": False,
                "enable_host_git_operations": False,
                "enable_session_store": False,
                "enable_skills": False,
                "skill_directories": [],
                "plugin_directories": [],
                "instruction_directories": [],
                "memory": {"enabled": False},
                "enable_mcp_apps": False,
            }
            agent = self.agent_factory(
                instructions=(
                    "You extract Shopify products only through the provided isolated tools. "
                    "Use the extraction tools directly or through execute_code to inspect evidence or workbook "
                    "ranges, validate, and call submit_extraction exactly once. Visual preview and structured "
                    "inspection are equally valid; choose what best fits the source. Never invent source locations."
                ),
                client=client,
                context_providers=[monty],
                tools=copilot_tools,
                default_options=options,
            )
            try:
                async with agent:
                    response = await agent.run(prompt)
                    usage = normalize_reported_token_usage(response)
                    if trace and usage is not None:
                        trace(
                            phase="copilot_usage",
                            message="Recorded Copilot BYOK extraction usage",
                            payload_preview={"usage_reported": True},
                            metadata={
                                "usage": usage,
                                "model_name": request.model.model_id,
                                "provider": request.model.provider,
                            },
                        )
            finally:
                await client.stop()
            if scoped.submitted is not None:
                result = scoped.submitted
            else:
                text = str(getattr(response, "text", "") or response)
                try:
                    result = ProductsList.model_validate_json(
                        _strip_markdown_json_fence(text)
                    )
                except ValidationError as exc:
                    raise RuntimeError(
                        "COPILOT_EXTRACTION_INVALID: agent did not submit a valid products payload"
                    ) from exc
            if trace:
                trace(
                    phase="copilot_session_done",
                    message="Completed isolated Copilot BYOK extraction",
                    payload_preview={
                        "products": len(result.products),
                        "preview_pages": len(preview.rendered_pages),
                    },
                )
            return result

    @staticmethod
    def _prompt(request: ExtractionRequest) -> str:
        source = (
            "canonical PDF evidence"
            if request.evidence_catalog is not None
            else "worksheet/CSV ranges"
        )
        correction = (
            f"\nValidation feedback from the server: {request.validation_feedback}"
            if request.validation_feedback
            else ""
        )
        return (
            f"Extract all Shopify products from {request.filename} using {source}. "
            "For PDF-derived values submit only canonical value_path and evidence_index in source_refs; "
            "never submit page, bbox, source_file_id, anchor_id, or provider. Each product owns its own "
            "source_refs, so value_path must be product-relative and must never start with products[<index>]. "
            "Use paths such as title, body_html, vendor, product_type, tags, and "
            "variant[<variant SKU or title>].sku or .price. Add a source_ref with a numeric evidence_index "
            "for every non-empty source-derived value; one evidence item may support multiple value paths. "
            "For values not present in the source, leave optional fields blank or explicitly mark their "
            "product-relative value_origins path as derived, defaulted, or ai_assisted. For spreadsheet values, "
            "include exact sheet and A1 cell/cell_range references. Inspect additional pages/ranges as needed. "
            "Validate the complete payload, correct it, then submit it." + correction
        )
