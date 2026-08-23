"""Reservation estimates and terminal settlement for paid processing."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
import os
from typing import Any

from application.services.document_formats import classify_document
from shared.credit_billing import (
    BillingRateUnavailable,
    GPT_5_6_LUNA_V1,
    ProviderUsage,
    calculate_cost,
    configured_platform_floors,
    estimate_usage,
)


def _document_dimensions(
    content: bytes, filename: str | None, content_type: str | None
) -> tuple[int, int]:
    document = classify_document(
        filename=filename, content_type=content_type, file_bytes=content
    )
    pages = 0
    rows = 0
    if document.kind == "pdf":
        try:
            import fitz

            with fitz.open(stream=content, filetype="pdf") as pdf:
                if pdf.needs_pass:
                    raise ValueError("Password-protected PDFs are not supported")
                pages = pdf.page_count
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Invalid or corrupt PDF") from exc
    elif document.kind == "image":
        pages = 1
    elif document.is_spreadsheet:
        try:
            if str(filename or "").lower().endswith(".csv"):
                rows = max(0, content.count(b"\n"))
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
                try:
                    rows = sum(max(0, sheet.max_row or 0) for sheet in workbook.worksheets)
                finally:
                    workbook.close()
        except Exception as exc:
            raise ValueError("Invalid or corrupt spreadsheet") from exc
    elif document.kind in {"docx", "pptx"}:
        # Exact converted-page count is unknown before the paid provider call;
        # reserve a conservative minimum and tune with benchmark P95 data.
        pages = 1
    return pages, rows


def quote_document_import(
    *,
    content: bytes,
    filename: str | None,
    content_type: str | None,
    model: str = "gpt-5.6-luna",
) -> tuple[Any, dict[str, int]]:
    pages, rows = _document_dimensions(content, filename, content_type)
    enforcement = os.getenv("CREDIT_ENFORCEMENT_ENABLED", "false").lower() == "true"
    if model != GPT_5_6_LUNA_V1.model and enforcement:
        raise BillingRateUnavailable(f"No effective billing rate for openai/{model}")
    azure_rate_raw = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_PAGE_RATE_USD", "").strip()
    azure_rate = Decimal(azure_rate_raw) if azure_rate_raw else (None if enforcement else Decimal("0"))
    floor_key = "BILLING_PLATFORM_FLOOR_DOCUMENT_IMPORT_USD"
    if enforcement and not os.getenv(floor_key, "").strip():
        raise BillingRateUnavailable("The calibrated document-import platform floor is not configured")
    usage = estimate_usage("document_import", pages=pages, rows=rows)
    quote = calculate_cost(
        usage=usage,
        openai_rate=GPT_5_6_LUNA_V1,
        azure_page_rate_usd=azure_rate,
        platform_cost_usd=configured_platform_floors(os.environ)["document_import"],
    )
    return quote, {"estimated_azure_pages": pages, "estimated_rows": rows}


def quote_intelligence_operation(
    operation_type: str, *, products: int, model: str = "gpt-5.6-luna"
) -> Any:
    enforcement = os.getenv("CREDIT_ENFORCEMENT_ENABLED", "false").lower() == "true"
    if model != GPT_5_6_LUNA_V1.model and enforcement:
        raise BillingRateUnavailable(f"No effective billing rate for openai/{model}")
    floor_key = f"BILLING_PLATFORM_FLOOR_{operation_type.upper()}_USD"
    if enforcement and not os.getenv(floor_key, "").strip():
        raise BillingRateUnavailable(f"The calibrated {operation_type} platform floor is not configured")
    return calculate_cost(
        usage=estimate_usage(operation_type, products=products),
        openai_rate=GPT_5_6_LUNA_V1,
        azure_page_rate_usd=None,
        platform_cost_usd=configured_platform_floors(os.environ).get(operation_type, Decimal("0")),
    )


def settle_terminal_run(service: Any, run: dict[str, Any]) -> dict[str, Any] | None:
    operation_id = str(run.get("billing_operation_id") or "").strip()
    if not operation_id:
        return None
    operation = service.get_credit_operation(operation_id)
    if not operation or operation.get("status") != "reserved":
        return operation

    prompt_tokens = max(0, int(run.get("prompt_tokens") or 0))
    completion_tokens = max(0, int(run.get("completion_tokens") or 0))
    provider_calls = service.list_provider_usage_events(operation_id)
    successful_paid_call = any(bool(item.get("succeeded", True)) for item in provider_calls)
    status = str(run.get("status") or "").lower()
    failure_class = str(run.get("billing_failure_class") or "").lower()
    metadata = operation.get("metadata") if isinstance(operation.get("metadata"), dict) else {}
    reported_azure_pages = 0
    for event in provider_calls:
        event_usage = event.get("usage") if isinstance(event.get("usage"), dict) else {}
        reported_azure_pages += max(0, int(event_usage.get("azure_pages") or 0))
    azure_pages = reported_azure_pages or max(0, int(metadata.get("estimated_azure_pages") or 0))

    if not successful_paid_call and not (prompt_tokens or completion_tokens):
        return service.release_credits(operation_id, "no_successful_paid_provider_call")
    if status in {"failed", "cancelled"} and failure_class != "merchant_after_paid":
        return service.release_credits(operation_id, "provider_or_platform_failure")

    missing_telemetry = successful_paid_call and not (prompt_tokens or completion_tokens)
    usage = ProviderUsage(
        input_tokens=prompt_tokens,
        cached_input_tokens=None,
        output_tokens=completion_tokens,
        azure_pages=azure_pages if successful_paid_call or status == "succeeded" else 0,
    )
    azure_rate_raw = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_PAGE_RATE_USD", "").strip()
    try:
        cost = calculate_cost(
            usage=usage,
            openai_rate=GPT_5_6_LUNA_V1,
            azure_page_rate_usd=(Decimal(azure_rate_raw) if azure_rate_raw else None),
            platform_cost_usd=configured_platform_floors(os.environ).get(
                str(operation.get("operation_type") or ""), Decimal("0")
            ),
            rate_card_version=str(operation.get("rate_card_version") or GPT_5_6_LUNA_V1.version),
        )
    except BillingRateUnavailable:
        # A successful call with unexpectedly incomplete telemetry/rates settles
        # at the reservation maximum; no post-provider work is made free.
        missing_telemetry = True
        return service.settle_credits(
            operation_id=operation_id,
            actual_credits=None,
            raw_usage=usage.as_dict(),
            cost_breakdown={"rate_unavailable_at_settlement": True},
            missing_telemetry=True,
        )
    return service.settle_credits(
        operation_id=operation_id,
        actual_credits=cost.credits,
        raw_usage=usage.as_dict(),
        cost_breakdown=cost.as_dict(),
        missing_telemetry=missing_telemetry,
    )
