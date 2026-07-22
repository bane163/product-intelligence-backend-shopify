import io

import openpyxl
import pytest

from ai.excel_reader import WorkbookValidationError, read_generated_workbook
from ai.excel_writer import create_excel_bytes
from ai.models import ProductInput, ProductOption, ProductVariant, ProductsList


def _generated() -> bytes:
    return create_excel_bytes(ProductsList(products=[
        ProductInput(
            handle="linen-shirt",
            title="Linen shirt",
            vendor="Stockpile",
            tags=["linen", "summer"],
            options=[ProductOption(name="Size", values=["S", "M"])],
            variants=[
                ProductVariant(option1="S", sku="LS-S", price=29.5, taxable=True),
                ProductVariant(option1="M", sku="LS-M", price=31, taxable=False),
            ],
            value_origins={"title": "source", "vendor": "merchant_edited"},
        )
    ]))


def test_generated_workbook_round_trips_supported_values_and_provenance():
    products = read_generated_workbook(_generated())
    assert len(products) == 1
    product = products[0]
    assert product["handle"] == "linen-shirt"
    assert product["title"] == "Linen shirt"
    assert product["tags"] == ["linen", "summer"]
    assert product["options"] == [{"name": "Size", "values": ["S", "M"]}]
    assert product["variants"][0]["sku"] == "LS-S"
    assert product["variants"][0]["price"] == 29.5
    assert product["variants"][1]["taxable"] is False
    assert product["value_origins"] == {"title": "source", "vendor": "merchant_edited"}


def test_reader_rejects_added_deleted_and_reordered_rows_with_cells():
    workbook = openpyxl.load_workbook(io.BytesIO(_generated()))
    sheet = workbook["Products"]
    sheet.move_range("A2:AY3", rows=1)
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(WorkbookValidationError) as exc:
        read_generated_workbook(output.getvalue())
    assert exc.value.errors
    assert exc.value.errors[0]["cell"].startswith("Products!")


def test_reader_rejects_variant_reordering_but_allows_value_edits():
    workbook = openpyxl.load_workbook(io.BytesIO(_generated()))
    sheet = workbook["Products"]
    price_column = next(cell.column for cell in sheet[1] if cell.value == "Variant Price")
    sheet.cell(2, price_column).value = 35
    output = io.BytesIO()
    workbook.save(output)
    assert read_generated_workbook(output.getvalue())[0]["variants"][0]["price"] == 35

    workbook = openpyxl.load_workbook(io.BytesIO(_generated()))
    sheet = workbook["Products"]
    first = [cell.value for cell in sheet[2]]
    second = [cell.value for cell in sheet[3]]
    for column, value in enumerate(second, start=1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(first, start=1):
        sheet.cell(3, column).value = value
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(WorkbookValidationError) as exc:
        read_generated_workbook(output.getvalue())
    assert any("reordered" in error["message"] for error in exc.value.errors)
